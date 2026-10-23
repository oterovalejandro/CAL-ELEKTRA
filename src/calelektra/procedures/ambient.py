"""Ambient conditions of a test from the TPH probe (temperature, pressure, humidity).

The probe dumps its readings to an SQL database and the laboratory exports the
query to an Excel workbook: sheet «DATOS» with the columns DATE, T, TUNIT, P,
PUNIT, H, HUNIT (one row per reading) and sheet «CABECERA» with the query.

The conditions of a test are the readings between its start and its end, plus
the last reading at or before the start and the first one at or after the end,
so a test shorter than the interval between readings is still covered. The workbook must
have readings no further than ``MAX_GAP`` from both ends of the test.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any

MAX_GAP = timedelta(minutes=15)
HEADER_ALIASES = {
    "date": {"DATE", "FECHA", "FECHA/HORA", "DATETIME"},
    "t": {"T", "TEMP", "TEMPERATURA"},
    "tunit": {"TUNIT"},
    "p": {"P", "PRESION", "PRESIÓN"},
    "punit": {"PUNIT"},
    "h": {"H", "HR", "HUMEDAD"},
    "hunit": {"HUNIT"},
}
DATE_FORMATS = ("%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M")


class AmbientError(ValueError):
    """The workbook cannot be read or does not cover the test."""


@dataclass(frozen=True)
class AmbientReading:
    time: datetime
    temperature_c: float
    pressure_hpa: float | None = None
    humidity_pct: float | None = None


def _number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(str(value).replace(",", "."))
    except ValueError:
        return None


def _date(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, str):
        for fmt in DATE_FORMATS:
            try:
                return datetime.strptime(value.strip(), fmt)
            except ValueError:
                continue
    return None


def _celsius(value: float, unit: Any) -> float:
    unit = str(unit or "C").strip().upper().replace("°", "").replace("º", "")
    if unit in ("C", ""):
        return value
    if unit == "K":
        return value - 273.15
    if unit == "F":
        return (value - 32) * 5 / 9
    raise AmbientError(f"Unidad de temperatura desconocida: «{unit}»")


def _hectopascal(value: float | None, unit: Any) -> float | None:
    if value is None:
        return None
    unit = str(unit or "Pa").strip().lower()
    factors = {"pa": 0.01, "hpa": 1.0, "mbar": 1.0, "kpa": 10.0, "bar": 1000.0}
    if unit not in factors:
        raise AmbientError(f"Unidad de presión desconocida: «{unit}»")
    return value * factors[unit]


def _stats(values: list[float], digits: int) -> dict[str, float] | None:
    if not values:
        return None
    return {"min": round(min(values), digits), "max": round(max(values), digits), "media": round(sum(values) / len(values), digits)}


@dataclass
class AmbientLog:
    """Readings of a probe workbook, sorted by time."""

    readings: list[AmbientReading]
    source: Path
    query: str | None = None

    @property
    def start(self) -> datetime:
        return self.readings[0].time

    @property
    def end(self) -> datetime:
        return self.readings[-1].time

    def describe(self) -> str:
        return f"{len(self.readings)} lecturas del {self.start:%d/%m/%Y %H:%M} al {self.end:%d/%m/%Y %H:%M}"

    def readings_for(self, start: datetime, end: datetime) -> list[AmbientReading]:
        """Readings that cover ``start`` – ``end``, including the ones just before and just after."""
        if end < start:
            start, end = end, start
        inside = [r for r in self.readings if start <= r.time <= end]
        before = next((r for r in reversed(self.readings) if r.time <= start), None)
        after = next((r for r in self.readings if r.time >= end), None)
        lower = before if before is not None and start - before.time <= MAX_GAP else None
        if lower is None and inside and inside[0].time - start <= MAX_GAP:
            lower = inside[0]
        upper = after if after is not None and after.time - end <= MAX_GAP else None
        if upper is None and inside and end - inside[-1].time <= MAX_GAP:
            upper = inside[-1]
        if lower is None or upper is None:
            raise AmbientError(
                f"El fichero de la sonda ({self.describe()}) no cubre el ensayo "
                f"({start:%d/%m/%Y %H:%M:%S} – {end:%H:%M:%S}): debe tener lecturas a menos de "
                f"{int(MAX_GAP.total_seconds() // 60)} min del inicio y del final"
            )
        return [r for r in self.readings if lower.time <= r.time <= upper.time]

    def conditions(self, start: datetime, end: datetime) -> dict[str, Any]:
        """Section ``condiciones`` of the record of a test made between ``start`` and ``end``."""
        readings = self.readings_for(start, end)
        section: dict[str, Any] = {
            "estado": "cargadas",
            "intervalo_ensayo": [f"{start:%Y-%m-%d %H:%M:%S}", f"{end:%Y-%m-%d %H:%M:%S}"],
            "lecturas": len(readings),
            "lecturas_desde": f"{readings[0].time:%Y-%m-%d %H:%M:%S}",
            "lecturas_hasta": f"{readings[-1].time:%Y-%m-%d %H:%M:%S}",
            "temperatura_c": _stats([r.temperature_c for r in readings], 2),
        }
        humidity = _stats([r.humidity_pct for r in readings if r.humidity_pct is not None], 1)
        pressure = _stats([r.pressure_hpa for r in readings if r.pressure_hpa is not None], 1)
        if humidity:
            section["humedad_pct"] = humidity
        if pressure:
            section["presion_hpa"] = pressure
        if self.query:
            section["consulta"] = self.query
        return section


def read_tph_workbook(path: str | Path) -> AmbientLog:
    """Readings of an Excel workbook exported from the TPH probe database."""
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - dependency of the package
        raise AmbientError("Falta openpyxl para leer ficheros .xlsx (pip install openpyxl)") from exc
    path = Path(path)
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several error types for files that are not workbooks
        raise AmbientError(f"No se puede leer {path.name}: {exc}") from exc
    try:
        query = None
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(min_row=1, max_row=5, max_col=2, values_only=True):
                if row and str(row[0] or "").strip().upper() == "CONSULTA" and len(row) > 1 and row[1]:
                    query = str(row[1]).strip()
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            columns: dict[str, int] = {}
            for _ in range(10):  # the header is in the first rows
                header = next(rows, None)
                if header is None:
                    break
                names = [str(cell).strip().upper() if cell is not None else "" for cell in header]
                columns = {key: names.index(alias) for key, aliases in HEADER_ALIASES.items() for alias in aliases if alias in names}
                if "date" in columns and "t" in columns:
                    break
            if "date" not in columns or "t" not in columns:
                continue
            readings = []
            for row in rows:
                values = {key: row[index] if index < len(row) else None for key, index in columns.items()}
                time, temperature = _date(values["date"]), _number(values["t"])
                if time is None or temperature is None:
                    continue
                readings.append(
                    AmbientReading(
                        time=time,
                        temperature_c=_celsius(temperature, values.get("tunit")),
                        pressure_hpa=_hectopascal(_number(values.get("p")), values.get("punit")),
                        humidity_pct=_number(values.get("h")),
                    )
                )
            if readings:
                readings.sort(key=lambda r: r.time)
                return AmbientLog(readings, path, query)
    finally:
        workbook.close()
    raise AmbientError(f"{path.name} no tiene lecturas de la sonda (hoja con las columnas DATE y T)")


def write_tph_workbook(path: str | Path, readings: list[AmbientReading], query: str = "Datos sonda TPH") -> Path:
    """Workbook with the layout of the laboratory export (demonstrations and tests)."""
    import openpyxl

    path = Path(path)
    workbook = openpyxl.Workbook()
    header = workbook.active
    header.title = "CABECERA"
    header.append(["CONSULTA", query])
    data = workbook.create_sheet("DATOS")
    data.append(["DATE", "T", "TUNIT", "P", "PUNIT", "H", "HUNIT"])
    for reading in readings:
        pressure = None if reading.pressure_hpa is None else round(reading.pressure_hpa * 100)
        data.append([reading.time, reading.temperature_c, "C", pressure, "Pa", reading.humidity_pct, "%"])
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path
