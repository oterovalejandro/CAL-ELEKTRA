"""Results of a cable by direction (VSWR, insertion loss, group delay and their uncertainty) and its certificate."""

from __future__ import annotations

import re
import struct
import zipfile
import zlib
from datetime import date
from pathlib import Path

import numpy as np
import pytest
from conftest import TODAY, tph_workbook
from docx import Document
from PySide6.QtWidgets import QApplication, QMessageBox

from calelektra.catalog.fields import FieldSpec
from calelektra.catalog.library import Library
from calelektra.core.bands import BandValues
from calelektra.core.frequency import FrequencyRange
from calelektra.core.networks import (
    group_delay,
    moving_average,
    smoothing_points,
    step_delay,
)
from calelektra.gui.certificate_dialog import CertificateDialog
from calelektra.gui.equipment_dialogs import EquipmentDialog
from calelektra.gui.field_form import FieldEditor
from calelektra.gui.main_window import MainWindow
from calelektra.gui.test_window import TestWindow
from calelektra.procedures.cable_results import (
    CableResults,
    report_frequencies,
    results_csv_path,
)
from calelektra.procedures.cable_specs import CableSpecification
from calelektra.procedures.engine import AutoOperator, ProcedureRunner, prepare_run
from calelektra.procedures.figures import (
    NAMES,
    delay_figure,
    insertion_loss_figure,
    range_of_use,
    vswr_figure,
)
from calelektra.procedures.model import ReportSettings
from calelektra.procedures.record_uncertainty import load_ambient_conditions
from calelektra.procedures.report import (
    ReportError,
    certificate_problem,
    comma,
    confidence_level,
    fixed,
    frequency_span,
    generate_certificate,
    prepare_certificate,
    stripped,
    vswr_note,
)
from calelektra.procedures.results import read_csv_columns
from calelektra.uncertainty.engine import UncertaintyEngine
from calelektra.uncertainty.traces import dataset_uncertainty

# the rows of the tables of the certificate of the laboratory, 10 MHz – 26.5 GHz
ROWS = [10e6, 1.3e9, 2.7e9, 4e9, 5.3e9, 6.6e9, 8e9, 9.3e9, 11e9, 12e9, 13e9, 15e9, 16e9, 17e9, 19e9, 20e9, 21e9, 23e9, 24e9, 25e9, 26.5e9]


def sucoflex_526s() -> CableSpecification:
    """The cable of the certificate of the laboratory: SF526S/11PC35/11PC35/2500mm, tolerances 1 dB and 5 %."""
    connector = BandValues.parse("1,12 hasta 18 GHz; 1,16 hasta 26,5 GHz")
    return CableSpecification(
        k1=0.229,
        k2=0.0071,
        frequency_unit="GHz",
        length_unit="1 m",
        length_m=2.5,
        vswr_a=connector,
        vswr_b=connector,
        insertion_loss_tolerance_db=1.0,
        vswr_tolerance_pct=5.0,
        nominal_range=FrequencyRange(0.0, 26.5e9),
    )


def line(f: np.ndarray, delay: float = 11.03e-9, loss_db: float = 1.0) -> np.ndarray:
    """S-parameters of a matched line: S11 = 0.04, S22 = 0.05, S21 = S12 with ``loss_db`` and ``delay``."""
    s = np.zeros((len(f), 2, 2), dtype=complex)
    s[:, 0, 0], s[:, 1, 1] = 0.04, 0.05j
    s[:, 1, 0] = s[:, 0, 1] = 10 ** (-loss_db / 20) * np.exp(-2j * np.pi * f * delay)
    return s


def png(path: Path, width: int = 40, height: int = 10) -> Path:
    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)

    rows = b"".join(b"\x00" + b"\x1f\x3a\x5f" * width for _ in range(height))
    header = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    path.write_bytes(b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", header) + chunk(b"IDAT", zlib.compress(rows)) + chunk(b"IEND", b""))
    return path


def test_group_delay_between_points_smoothed_as_the_laboratory():
    f = np.linspace(10e6, 26.5e9, 2650)
    s21 = 0.9 * np.exp(-2j * np.pi * f * 11.03e-9)
    assert np.allclose(step_delay(f, s21), 11.03e-9, atol=1e-15) and np.allclose(group_delay(f, s21), 11.03e-9, atol=1e-15)
    # a list with steps where a long cable turns more than half a turn: unwrapped with the delay of the whole trace
    listed = np.concatenate([np.arange(30e6, 1000e6, 10e6), np.arange(1000e6, 18000e6 + 1, 100e6)])
    assert np.allclose(group_delay(listed, 0.8 * np.exp(-2j * np.pi * listed * 40e-9 + 0.4j)), 40e-9, rtol=1e-9)
    gaps = s21.copy()
    gaps[100:110] = np.nan
    delay = group_delay(f, gaps)
    assert np.isnan(delay[100:110]).all() and np.allclose(delay[~np.isnan(delay)], 11.03e-9, atol=1e-15)
    # the mean of the delays between points over 1 % of the points: 26 of 2650, a ripple of 2 points averages out
    rippled = s21 * np.exp(1j * np.where(np.arange(len(f)) % 2, 0.01, -0.01))
    assert smoothing_points(len(f)) == 26 and np.allclose(group_delay(f, rippled)[26:-26], 11.03e-9, rtol=1e-6, atol=0)
    assert not np.allclose(step_delay(f, rippled), 11.03e-9, rtol=1e-3, atol=0)
    assert smoothing_points(150) == 1 and np.array_equal(group_delay(f[:150], rippled[:150]), step_delay(f[:150], rippled[:150]))


def test_results_by_direction_with_their_uncertainty_and_limits(tmp_path):
    f = np.linspace(10e6, 26.5e9, 2650)
    uncertainty = {p: (np.full(len(f), 0.5), np.full(len(f), 3.0)) for p in ("S11", "S22")}
    uncertainty |= {p: (np.where(f <= 10e9, 0.04, 0.05), np.full(len(f), 0.4)) for p in ("S21", "S12")}
    results = CableResults.compute(f, line(f), uncertainty, sucoflex_526s())
    fwd, rev = results.directions["Fwd"], results.directions["Rev"]
    assert (fwd.reflection, fwd.transmission, fwd.port) == ("S11", "S21", 1) and (rev.reflection, rev.transmission, rev.port) == ("S22", "S12", 2)
    assert np.allclose(fwd.vswr, 1.04 / 0.96) and np.allclose(rev.vswr, 1.05 / 0.95)
    # U_VSWR = 2·U_ρ/(1 − ρ)², U_ρ = ρ·ln10/20·U_dB
    assert np.allclose(fwd.vswr_u, 2 * 0.04 * np.log(10) / 20 * 0.5 / 0.96**2)
    # insertion loss and its uncertainty: those of |S21| in dB
    assert np.allclose(fwd.insertion_loss_db, -1.0) and np.allclose(fwd.insertion_loss_u, uncertainty["S21"][0])
    # group delay and U_τ = U_φ/(360°·f): 0.1111 ns at 10 MHz with 0.4°, as in the certificate
    assert np.allclose(fwd.group_delay_ns, 11.03, atol=1e-6) and fixed(fwd.group_delay_u_ns[0], "ns") == "0.1111 ns"
    # limits of the certificate of the laboratory
    first, last = 0, len(f) - 1
    assert fixed(fwd.insertion_loss_limit[first], "dB") == "-1.0574 dB" and fixed(fwd.insertion_loss_limit[last], "dB") == "-4.4175 dB"
    # VSWR: 5 % on the assembled cable (1.2553 and 1.3478 up to 18 and 26.5 GHz)
    assert stripped(fwd.vswr_limit[first]) == "1.3181" and stripped(fwd.vswr_limit[last]) == "1.4152"
    assert fwd.passes().all() and rev.passes().all()
    summary = results.summary()
    assert (summary["suavizado_pct"], summary["suavizado_puntos"]) == (1.0, 26) and summary["Fwd"]["retardo_grupo_ns"]["media"] == pytest.approx(11.03)
    assert summary["Rev"]["vswr_max"] == pytest.approx(1.1053, abs=1e-4) and summary["Fwd"]["margen_perdida_insercion_db"] > 0
    header, columns = read_csv_columns(results.write_csv(results_csv_path(tmp_path / "DUT.s2p")))
    assert header[:4] == ["Frecuencia (Hz)", "VSWR puerto 1 (Fwd)", "VSWR puerto 1 (Fwd, suavizado 1 %)", "U VSWR puerto 1 (Fwd, k=2)"]
    assert "Retardo de grupo S12 (Rev, ns, suavizado 1 %)" in header and columns.shape[1] == len(f)


def test_vswr_uncertainty_of_each_port_at_its_own_level(library):
    """U_VSWR = 2·ρ·ln10/20·U_dB/(1 − ρ)²; in Rev, U_dB is that of S22 at the level of S22 (not that of S11)."""
    f = np.array([1e9, 4e9, 5.3e9, 12e9, 23e9])
    s = line(f)
    s[:, 0, 0], s[:, 1, 1] = 0.04, 0.1  # S11 at −28 dB, S22 at −20 dB: two ranges of the uncertainty of the ZNB43
    engine = UncertaintyEngine(library.uncertainty_path(library.equipment["VNA-01"]))
    traces = dataset_uncertainty(engine, "ZNB43", f, s)
    u11, u22 = traces["S11"].magnitude_db, traces["S22"].magnitude_db
    assert (u11 > 2 * u22).all()  # the level decides the range: far apart
    results = CableResults.compute(f, s, traces)
    fwd, rev = results.directions["Fwd"], results.directions["Rev"]
    assert np.allclose(fwd.vswr_u, 2 * 0.04 * np.log(10) / 20 * u11 / 0.96**2)
    assert np.allclose(rev.vswr_u, 2 * 0.1 * np.log(10) / 20 * u22 / 0.9**2)  # the factor 2 of dVSWR/dρ = 2/(1 − ρ)²
    with np.errstate(divide="ignore"):  # the same VSWR uncertainty with the derivative of VSWR(ρ)
        derivative = (results.directions["Rev"].vswr - (1 + 0.1 * 10 ** (-1e-6 / 20)) / (1 - 0.1 * 10 ** (-1e-6 / 20))) / 1e-6
    assert np.allclose(rev.vswr_u, derivative * u22, rtol=1e-4)


def test_charts_of_the_certificate_in_the_format_of_the_laboratory():
    f = np.linspace(10e6, 26.5e9, 2650)
    uncertainty = {p: (np.full(len(f), 0.5), np.full(len(f), 3.0)) for p in ("S11", "S22")}
    uncertainty |= {p: (np.full(len(f), 0.04), np.full(len(f), 0.4)) for p in ("S21", "S12")}
    results = CableResults.compute(f, line(f), uncertainty, sucoflex_526s())
    assert range_of_use(results).all()  # the limits cover the sweep: the axes, over it
    rev = results.directions["Rev"]
    figure = vswr_figure(results, rev)
    (axes,) = figure.axes
    assert tuple(figure.get_size_inches()) == (16, 4) and (axes.get_ylabel(), axes.get_xlabel()) == ("VSWR Port 2", "Frec (Hz)")
    assert [text.get_text() for text in axes.get_legend().get_texts()] == ["S22 (Reflection)", f"u_exp ±{np.mean(rev.vswr_u):.3f}", "Limit"]
    assert axes.get_xlim() == (f[0], f[-1]) and axes.xaxis.get_major_formatter()(5e9, 0) == "5.0e+09"
    assert axes.get_ylim() == pytest.approx((1.05 / 0.95 - 0.1, 1.05 / 0.95 + 0.3))  # a flat VSWR: the margins only
    loss = insertion_loss_figure(results, results.directions["Fwd"])
    assert [item.get_ylabel() for item in loss.axes] == ["S21 (dB)", "S21 (º)"]
    assert [text.get_text() for text in loss.axes[0].get_legend().get_texts()] == ["S21 (Insertion Loss)", "u_exp(dB) ±0.040", "Limit"]
    assert [text.get_text() for text in loss.axes[1].get_legend().get_texts()] == ["S21 (Insertion Loss)", "u_exp(º) ±0.400"]
    delay = delay_figure(results, results.directions["Fwd"]).axes[0]
    assert delay.get_ylabel() == "S21 (ns)" and delay.get_ylim() == pytest.approx((11.03 * 0.7, 11.03 * 1.3), rel=1e-6)
    assert [text.get_text() for text in delay.get_legend().get_texts()][0] == "S21 (Transmision)"


def test_rows_of_the_certificate_at_the_nearest_measured_point():
    assert report_frequencies(21, 10e6, 26.5e9) == ROWS
    assert report_frequencies(["1 GHz", "30 GHz", "500 MHz"], 10e6, 26.5e9) == [500e6, 1e9]  # a list: those inside
    # a logarithmic sweep: logarithmically spaced, two significant digits
    logarithmic = report_frequencies(21, 9e3, 3e9, logarithmic=True)
    assert len(logarithmic) == 21 and logarithmic[:4] == [9e3, 17e3, 32e3, 61e3] and logarithmic[-3:] == [840e6, 1.6e9, 3e9]
    f = np.linspace(10e6, 26.5e9, 101)  # steps of 264.9 MHz
    rows = CableResults.compute(f, line(f), None, sucoflex_526s()).report_rows(ROWS)
    assert len(rows["Fwd"]) == len(rows["Rev"]) == 21
    assert rows["Fwd"][1]["frecuencia"] == 1.3e9 and rows["Fwd"][1]["frecuencia_medida"] == pytest.approx(f[5])
    assert np.isnan(rows["Fwd"][1]["u_vswr"]) and rows["Rev"][0]["pasa"]  # without uncertainty


def test_formats_and_texts_of_the_certificate():
    assert comma(0.229) == "0,229" and comma(24.66, 1) == "24,7" and comma(None) == "—"
    assert stripped(1.0748) == "1.0748" and stripped(0.00298) == "0.003" and fixed(-0.063, "dB") == "-0.0630 dB"
    assert confidence_level(2) == "95,5%" and confidence_level(1) == "68,3%"
    assert frequency_span(10e6, 18e9) == "10 MHz y 18 GHz" and frequency_span(18e9, 26.5e9) == "18 y 26,5 GHz"
    # the footnote of the VSWR of the certificate of the laboratory, word for word
    connectors = (
        "Entre 10 MHz y 18 GHz, el VSWR máximo declarado de cada conector es 1,12:1. "
        "Entre 18 y 26,5 GHz, el VSWR máximo declarado de cada conector es 1,16:1. "
    )
    assert vswr_note(sucoflex_526s(), FrequencyRange(10e6, 26.5e9)) == (
        connectors + "El VSWR medido no deberá superar 1,32:1 entre 10 MHz y 18 GHz, ni 1,42:1 entre 18 y 26,5 GHz."
    )
    per_connector = sucoflex_526s()
    per_connector.vswr_tolerance_on = "conectores"  # the limits of the example certificate of the laboratory
    assert vswr_note(per_connector, FrequencyRange(10e6, 26.5e9)) == (
        connectors + "El VSWR medido no deberá superar 1,39:1 entre 10 MHz y 18 GHz, ni 1,49:1 entre 18 y 26,5 GHz."
    )
    other = sucoflex_526s()
    other.vswr_b = BandValues.single(1.2)
    assert "del conector A es 1,16:1 y el del conector B, 1,2:1" in vswr_note(other, FrequencyRange(20e9, 26.5e9))


def run_example(library: Library, tmp_path: Path):
    record = ProcedureRunner(library, library.procedures["cable_CAB-DUT-0006"], AutoOperator(), today=TODAY, seed=1).run().record
    load_ambient_conditions(library, record, tph_workbook(tmp_path / "tph.xlsx", record.started, record.finished))
    return record


def docx_xml(path: Path) -> dict[str, str]:
    with zipfile.ZipFile(path) as archive:
        return {name: archive.read(name).decode("utf-8") for name in archive.namelist() if name.endswith(".xml")}


def test_certificate_of_the_example_cable(library, tmp_path):
    record = run_example(library, tmp_path)
    assert record.verdict == "PASA" and record.results["sentidos"]["suavizado_pct"] == 1.0
    entry = next(e for e in record.equipment if e["id"] == "VNA-01")  # the record keeps what the certificate needs
    assert entry["referencia"] == "LE1941.000" and entry["fabricante"] == "Rohde & Schwarz" and entry["designacion"] == "ZNB43"
    assert record.dut["referencia"] == "LE2333.001" and record.dut["denominacion"] == "Cable Sucoflex 526s 3.5mm-3.5mm"

    path, certificate = generate_certificate(library, record, report_date=date(2026, 10, 27), figures={"retardo": png(tmp_path / "otra.png")})
    reference = f"LE2333.001_{record.started:%Y_%m}"
    assert path == record.folder / "certificado" / f"{reference}.docx" and path.is_file()
    context = certificate.context
    assert context["eut"] == {
        "id": "CAB-DUT-0006",
        "referencia": "LE2333.001",
        "instrumento": "Cable Sucoflex 526s 3.5mm-3.5mm",
        "fabricante": "Huber+Suhner",
        "modelo": "SF526S/11PC35/11PC35/2500mm",
        "serie": "682138001",
    }
    assert [row["frecuencia"] for row in context["fwd"]] == [
        "10 MHz", "1.3 GHz", "2.7 GHz", "4 GHz", "5.3 GHz", "6.6 GHz", "8 GHz", "9.3 GHz", "11 GHz", "12 GHz", "13 GHz",
        "15 GHz", "16 GHz", "17 GHz", "19 GHz", "20 GHz", "21 GHz", "23 GHz", "24 GHz", "25 GHz", "26.5 GHz",
    ]  # fmt: skip
    first, last = context["fwd"][0], context["rev"][-1]
    assert (first["limite_vswr"], first["limite_perdida"], first["u_retardo"]) == ("1.3181", "-1.0574 dB", "0.1111 ns")
    assert (last["limite_vswr"], last["limite_perdida"]) == ("1.4152", "-4.4175 dB")
    assert context["ambiente"] == {"temperatura": {"min": "23,0", "max": "23,0"}, "humedad": {"min": "37,0", "max": "37,2"}, "presion": {"min": "94260", "max": "94260"}}
    assert [row["inventario"] for row in context["equipos"]] == ["LE1941.000", "LE1941.002", "ADP-NMD292-01", "LE1941.001", "LE1689.000"]
    assert context["equipos"][-1]["proxima_calibracion"] == "20/12/2027"
    assert any("ADP-NMD292-01" in warning for warning in certificate.warnings)
    assert not any("ráfica" in warning for warning in certificate.warnings)
    # the charts, in the format of the laboratory, next to the certificate (the one given replaces the one drawn)
    figures = path.with_name(f"{reference}_figuras")
    assert sorted(item.name for item in figures.iterdir()) == sorted(NAMES.values())

    # the document: no tag left, the values where the laboratory has them
    parts = docx_xml(path)
    assert not [name for name, xml in parts.items() if "{{" in xml or "{%" in xml]
    document = Document(str(path))
    assert document.core_properties.title == reference
    header = next(xml for name, xml in parts.items() if name.startswith("word/header") and "LE2306.000" not in xml and "Ref.:" in xml)
    assert f"<w:t xml:space=\"preserve\">{reference}</w:t>" in header  # «Ref.»: the title, also in the document properties
    assert ">LE2333.001<" in header and ">27/10/2026<" in header and "DATE" not in header
    body = "".join(re.findall(r"<w:t(?: [^>]*)?>([^<]*)</w:t>", parts["word/document.xml"]))  # the cover is a content control
    assert "Bienal" in body and "ATN-Q-IO-22-164" in body and "k=2 (nivel de confianza aproximado del 95,5%)" in body
    assert "rango de frecuencias de 10 MHz a 26,5 GHz" in body and "tolerancia de 1 dB" in body and "tolerancia de un 5%" in body
    assert "[Gráfica" not in body and "Resultado final: PASA" in body
    tables = {table.rows[0].cells[0].text: table for table in document.tables}
    fwd, rev = [table for table in document.tables if table.rows[0].cells[0].text == "Frecuencia"]
    assert len(fwd.rows) == len(rev.rows) == 22 and fwd.rows[1].cells[0].text == "10 MHz" and rev.rows[21].cells[0].text == "26.5 GHz"
    assert rev.rows[0].cells[4].text == "Pérdida de inserción Rev (dB)"
    assert [row.cells[0].text for row in tables["N.º de inventario"].rows[1:]] == [row["inventario"] for row in context["equipos"]]
    assert "1,32:1 entre 10 MHz y 18 GHz, ni 1,42:1 entre 18 y 26,5 GHz" in parts["word/footnotes.xml"]
    assert "a = 0,229 dB/m y b = 0,0071 dB/m" in parts["word/footnotes.xml"]
    images = [r for r in document.part.rels.values() if r.reltype.endswith("/image")]
    assert len(images) == 5 and png(tmp_path / "otra.png").read_bytes() in [r.target_part.blob for r in images]
    saved = library.load_test(record.folder)
    assert saved.results["certificado"]["referencia"] == reference and f"certificado/{reference}.docx" in saved.files
    assert saved.results["certificado"]["figuras"] == f"certificado/{reference}_figuras" and saved.results["certificado"]["suavizado_pct"] == 1.0


def test_certificate_only_of_finished_two_port_tests(library, tmp_path):
    record = run_example(library, tmp_path)
    record.status = "abortado"
    assert certificate_problem(record) == "El ensayo no está completado (Abortado)"
    with pytest.raises(ReportError):
        prepare_certificate(library, record)
    record.status = "completado"
    record.procedure["algoritmo"] = "vna_1_puerto"
    assert "2 puertos" in certificate_problem(record)
    record.procedure["algoritmo"] = "vna_2_puertos"
    (library.root / "informes" / "certificado_cable.docx").unlink()
    with pytest.raises(ReportError, match="No se encuentra la plantilla"):
        generate_certificate(library, record)


@pytest.fixture(scope="module")
def qapp():
    return QApplication.instance() or QApplication([])


def test_certificate_from_the_test_window(qapp, library, tmp_path):
    record = run_example(library, tmp_path)
    window = TestWindow(library, record=record, today=TODAY)
    assert window.certificate_button.isEnabled() and window.action_certificate.isEnabled()
    target = tmp_path / "certificado.docx"
    assert window.make_certificate(target) == target and target.is_file()
    boxes = [box for box in window.findChildren(QMessageBox) if box.isVisible()]
    assert boxes and boxes[0].text().startswith(f"Certificado guardado en {target}")
    assert "Abrir" in [button.text() for button in boxes[0].buttons()]
    for box in boxes:
        box.close()
    other = TestWindow(library, library.procedures["cable_CAB-DUT-0002"], library.tests_folder, "Cable", "prueba", today=TODAY)
    assert not other.certificate_button.isEnabled()  # nothing measured yet
    for test in (window, other):
        test.close()
        test.deleteLater()


def test_tables_fwd_and_rev_in_the_test_window(qapp, library, tmp_path):
    record = run_example(library, tmp_path)
    window = TestWindow(library, record=record, today=TODAY)
    tabs = [window.tabs.tabText(i) for i in range(window.tabs.count())]
    assert tabs[:3] == ["Gráfica", "Tabla", "Pasa / falla — PASA"] and window.tabs.isTabVisible(1)
    view = window.report_tables
    fwd, rev = view.tables["fwd"], view.tables["rev"]
    assert fwd.rowCount() == rev.rowCount() == 21  # the frequencies of the certificate
    context = prepare_certificate(library, record).context
    first = [fwd.item(0, column).text() for column in range(fwd.columnCount())]
    expected = context["fwd"][0]
    keys = ("frecuencia", "vswr", "u_vswr", "limite_vswr", "perdida", "u_perdida", "limite_perdida", "retardo", "u_retardo")
    assert first == [expected[key] for key in keys] and rev.item(20, 0).text() == "26.5 GHz"
    assert "21 frecuencias igualmente espaciadas de 10 MHz a 26,5 GHz" in view.note.text()
    text = view.text()  # «Copiar tablas»: tab-separated, the headers on one line
    assert text.startswith("Fwd\nFrecuencia\tVSWR puerto 1\tIncertidumbre VSWR") and "\n\nRev\n" in text
    other = TestWindow(library, library.procedures["lisn_LISN-0001"], library.tests_folder, "LISN", "prueba", today=TODAY)
    assert not other.tabs.isTabVisible(other.tabs.indexOf(other.report_tables))  # only in the 2-port tests
    for test in (window, other):
        test.close()
        test.deleteLater()


def test_inventory_reference_and_name_of_an_equipment(qapp, library):
    dialog = EquipmentDialog(library, library.equipment["VNA-01"])
    assert dialog.reference.text() == "LE1941.000" and dialog.name_edit.text() == "Analizador vectorial de redes"
    dialog.reference.setText("LE1941.010")
    equipment = dialog.build()
    assert equipment.reference == "LE1941.010" and equipment.name == "Analizador vectorial de redes"
    assert equipment.to_yaml_dict()["referencia"] == "LE1941.010"
    copy = EquipmentDialog(library, copy_of=library.equipment["VNA-01"])
    assert copy.reference.text() == "" and copy.name_edit.text() == "Analizador vectorial de redes"  # a copy is another unit
    for widget in (dialog, copy):
        widget.close()
        widget.deleteLater()


def test_certificate_settings_in_the_template_editor(qapp, library):
    window = MainWindow(library)
    window.show_section("procedimientos")
    view = window.procedures_view
    view.reload(select="cable_CAB-DUT-0006")
    editor = view.editor
    assert editor.report_card.isVisibleTo(editor) and editor.report_sensor.currentData() == "TPH-01"
    assert (editor.report_procedure.text(), editor.report_interval.text(), editor.report_rows.text()) == ("ATN-Q-IO-22-164", "Bienal", "21")
    assert editor.build().report == library.procedures["cable_CAB-DUT-0006"].report  # saving from the editor keeps it
    assert "1 %" in editor.report_smoothing.text()  # always as the laboratory: not a setting
    editor.report_rows.setText("10 MHz; 1 GHz; 26,5 GHz")
    assert editor.build().report.frequencies == ["10 MHz", "1 GHz", "26,5 GHz"]
    # the settings of 0.10.1 are read and dropped
    assert ReportSettings.model_validate({"apertura_retardo": 10, "suavizado_pct": 1}) == ReportSettings()
    editor.report_rows.setText("10 MHz; uno")
    with pytest.raises(ValueError, match="no es una frecuencia"):
        editor.build()
    view._set_dirty(False)
    view.reload(select="lisn_LISN-0001")
    assert not editor.report_card.isVisibleTo(editor)
    window.close()


def test_bands_entered_in_a_table(qapp):
    spec = FieldSpec(label="VSWR del conector A", type="numero_bandas", minimo=1)
    editor = FieldEditor(spec)
    editor.set_value([{"hasta": "18 GHz", "valor": 1.12}, {"hasta": "26.5 GHz", "valor": 1.16}])
    dialog = editor.edit_bands()  # the button «Bandas…»
    assert dialog.rows() == [{"hasta": "18 GHz", "valor": "1,12"}, {"hasta": "26,5 GHz", "valor": "1,16"}]
    dialog.add_band("40 GHz", "1,3")
    assert dialog.table.item(2, 0).text() == "26,5 GHz"  # each band starts where the previous one ends
    dialog.accept()
    assert editor.edit.text() == "1,12 hasta 18 GHz; 1,16 hasta 26,5 GHz; 1,3 hasta 40 GHz"
    assert editor.value().at([30e9]).tolist() == [1.3]
    wrong = editor.edit_bands()
    wrong.table.item(0, 2).setText("0,9")
    wrong.accept()  # not accepted: the error is shown in the dialog
    assert wrong.isVisible() and "mínimo" in wrong.error.text()
    wrong.reject()
    for row in range(wrong.table.rowCount()):  # an empty table clears the field
        wrong.table.item(row, 1).setText("")
        wrong.table.item(row, 2).setText("")
    assert wrong.bands() is None


def test_certificate_dialog_loads_the_workbook_of_the_probe(qapp, library, tmp_path):
    record = ProcedureRunner(library, library.procedures["cable_CAB-DUT-0006"], AutoOperator(), today=TODAY, seed=1).run().record
    dialog = CertificateDialog(library, record)
    assert dialog.conditions_button.isVisibleTo(dialog) and "condiciones ambientales" in dialog.warnings.text()
    assert dialog.load_conditions(tph_workbook(tmp_path / "tph.xlsx", record.started, record.finished))
    for box in dialog.findChildren(QMessageBox):
        box.close()
    assert not dialog.conditions_button.isVisibleTo(dialog) and record.uncertainty["estado"] == "completa"
    assert dialog.certificate.context["ambiente"]["temperatura"] == {"min": "23,0", "max": "23,0"}
    dialog.close()


def test_vswr_tolerance_choice_in_the_template_editor(qapp, library):
    window = MainWindow(library)
    window.show_section("procedimientos")
    view = window.procedures_view
    view.reload(select="cable_CAB-DUT-0006")
    editor = view.editor
    assert editor.vswr_on.currentData() == "cable" and editor.build().cable_tolerances.vswr_on == "cable"
    assert "vswr_sobre" not in editor.build().to_yaml_dict()["tolerancias_cable"]  # the default is not written
    editor.vswr_on.setCurrentIndex(editor.vswr_on.findData("conectores"))
    built = editor.build()
    assert built.cable_tolerances.vswr_on == "conectores" and built.to_yaml_dict()["tolerancias_cable"]["vswr_sobre"] == "conectores"
    plan = prepare_run(library, built, TODAY)
    assert stripped(plan.cable_spec.max_vswr(np.array([10e6]))[0]) == "1.386"  # the example certificate of the laboratory
    view._set_dirty(False)
    window.close()


def test_smoothing_of_the_tables_as_the_certificates_of_the_lab():
    """Always: the VSWR and the group delay over 1 % of the points; the insertion loss and the verdict, of the point."""
    # moving average with the ends mirrored, as scipy.ndimage.uniform_filter1d(x, size)
    x = np.arange(10, dtype=float)
    smooth = moving_average(x, 4)  # window [k − 2, k + 1]
    assert smooth[0] == pytest.approx((1 + 0 + 0 + 1) / 4) and smooth[5] == pytest.approx(4.5) and smooth[9] == pytest.approx((7 + 8 + 9 + 9) / 4)
    assert np.array_equal(moving_average(x, 1), x)
    f = np.linspace(10e6, 26.5e9, 5299)
    assert np.allclose(step_delay(f, np.exp(-2j * np.pi * f * 11.03e-9)), 11.03e-9)
    s = line(f)
    s[:, 0, 0] = np.where(np.arange(len(f)) % 2, 0.02, 0.06)  # a VSWR that jumps from point to point
    results = CableResults.compute(f, s, None, sucoflex_526s())
    assert results.smoothing == 52  # 1 % of 5299 points, as the laboratory
    rows = results.report_rows(ROWS)
    middle = 10  # 13 GHz, far from the ends: half the points of each value
    index = int(np.argmin(np.abs(f - ROWS[middle])))
    fwd = results.directions["Fwd"]
    assert rows["Fwd"][middle]["vswr"] == pytest.approx(((1.02 / 0.98) + (1.06 / 0.94)) / 2)
    assert fwd.vswr[index] in (pytest.approx(1.02 / 0.98), pytest.approx(1.06 / 0.94))  # that of the point: the verdict
    assert rows["Fwd"][middle]["perdida_insercion"] == fwd.insertion_loss_db[index]  # not smoothed
    assert rows["Fwd"][middle]["retardo_grupo"] == pytest.approx(11.03, abs=1e-6)
